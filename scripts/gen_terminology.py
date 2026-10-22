#!/usr/bin/env python3
"""由「附錄四_通報代碼彙整表」與「附錄四_送驗代碼彙整表」(xlsx) 產生 FSH 代碼系統、值集與概念對應。

用法：
    python3 scripts/gen_terminology.py [來源目錄] [輸出目錄]

預設來源目錄為 source/nidrs-emr2.0-v1.42，輸出目錄為 input/fsh/generated。
產出的 .fsh 檔請勿手動修改；代碼表改版時重新執行本程式即可。
"""
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

SRC = Path(sys.argv[1] if len(sys.argv) > 1 else "source/nidrs-emr2.0-v1.42")
OUT = Path(sys.argv[2] if len(sys.argv) > 2 else "input/fsh/generated")
REPORT_XLSX = SRC / "附錄四_通報代碼彙整表_v1.42.xlsx"
SAMPLE_XLSX = SRC / "附錄四_送驗代碼彙整表_v1.42.xlsx"
CODE_VERSION = "v1.42"
CODE_DATE = "2026-09-29"

INACTIVE_MARKS = ("停用", "刪除", "移除")
SAFE_CODE = re.compile(r"^[A-Za-z0-9][A-Za-z0-9._]*$")


def s(v):
    """將儲存格內容轉為去頭尾空白之字串；整數型態的代碼不帶小數點。"""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def fsh_str(text):
    text = re.sub(r"\s+", " ", text).strip()
    return text.replace("\\", "\\\\").replace('"', '\\"')


def fsh_code(code):
    return "#" + code if SAFE_CODE.match(code) else '#"' + fsh_str(code) + '"'


def rows(wb, sheet, skip=1):
    for r in list(wb[sheet].iter_rows(values_only=True))[skip:]:
        yield [s(c) for c in r]


class CS:
    def __init__(self, cid, name, title, sheet, description="", props=None, hierarchical=False):
        self.cid, self.name, self.title, self.sheet = cid, name, title, sheet
        self.description = description
        self.props = props or OrderedDict()  # code -> (uri, type, description)
        self.concepts = OrderedDict()  # code -> dict(display, inactive, comment, props, parent)
        self.hierarchical = hierarchical
        self.dups = []

    def add(self, code, display, remark="", parent=None, **props):
        if not code or not display:
            return
        inactive = any(m in remark for m in INACTIVE_MARKS)
        if code in self.concepts:
            if self.concepts[code]["display"] == display:
                return
            self.dups.append(code)
            # 修剪空白後代碼相同時，以仍在使用中之代碼為準
            if not (self.concepts[code]["inactive"] and not inactive):
                return
        self.concepts[code] = dict(
            display=display,
            inactive=inactive,
            comment=remark,
            parent=parent,
            props={k: v for k, v in props.items() if v},
        )

    def fsh(self):
        desc = (
            f"臺灣衛生福利部疾病管制署 NIDRS 電子病歷傳染病自動通報作業 2.0「{self.sheet}」代碼，"
            f"資料來源：附錄四代碼彙整表 {CODE_VERSION}（{CODE_DATE}）；資料所屬單位：臺灣衛生福利部疾病管制署。"
        )
        if self.description:
            desc += self.description
        out = [
            f"CodeSystem: {self.name}",
            f"Id: {self.cid}",
            f'Title: "{fsh_str(self.title)}"',
            f'Description: "{fsh_str(desc)}"',
            "* ^caseSensitive = true",
            "* ^content = #complete",
            "* ^experimental = false",
            f'* ^date = "{CODE_DATE}"',
            '* ^version = "0.1.0"',
        ]
        if self.hierarchical:
            out.append("* ^hierarchyMeaning = #is-a")
        props = OrderedDict(
            inactive=("http://hl7.org/fhir/concept-properties#inactive", "boolean", "停用、刪除或移除之代碼"),
            comment=("http://hl7.org/fhir/concept-properties#comment", "string", "代碼表備註(代碼處理方式)"),
        )
        props.update(self.props)
        used = {k for c in self.concepts.values() for k in c["props"]}
        if any(c["inactive"] for c in self.concepts.values()):
            used.add("inactive")
        if any(c["comment"] for c in self.concepts.values()):
            used.add("comment")
        if any(c["parent"] for c in self.concepts.values()):
            props["parent"] = ("http://hl7.org/fhir/concept-properties#parent", "code", "上層代碼")
            used.add("parent")
        i = 0
        for k, (uri, typ, d) in props.items():
            if k not in used:
                continue
            out += [
                f"* ^property[{i}].code = #{k}",
                f'* ^property[{i}].uri = "{uri}"',
                f'* ^property[{i}].description = "{fsh_str(d)}"',
                f"* ^property[{i}].type = #{typ}",
            ]
            i += 1
        for code, c in self.concepts.items():
            out.append(f'* {fsh_code(code)} "{fsh_str(c["display"])}"')
            j = 0
            if c["inactive"]:
                out += [f"  * ^property[{j}].code = #inactive", f"  * ^property[{j}].valueBoolean = true"]
                j += 1
            if c["comment"]:
                out += [f"  * ^property[{j}].code = #comment", f'  * ^property[{j}].valueString = "{fsh_str(c["comment"])}"']
                j += 1
            if c["parent"]:
                out += [f"  * ^property[{j}].code = #parent", f"  * ^property[{j}].valueCode = {fsh_code(c['parent'])}"]
                j += 1
            for k, v in c["props"].items():
                typ = props[k][1]
                if typ == "code":
                    out += [f"  * ^property[{j}].code = #{k}", f"  * ^property[{j}].valueCode = {fsh_code(v)}"]
                else:
                    out += [f"  * ^property[{j}].code = #{k}", f'  * ^property[{j}].valueString = "{fsh_str(v)}"']
                j += 1
        out += [
            "",
            f"ValueSet: {self.name}",
            f"Id: {self.cid}",
            f'Title: "{fsh_str(self.title)}值集"',
            f'Description: "{fsh_str(self.title)}值集，包含代碼系統 {self.name} 之所有代碼。"',
            "* ^experimental = false",
            f'* ^date = "{CODE_DATE}"',
            '* ^version = "0.1.0"',
            f"* include codes from system {self.name}",
            "",
        ]
        return "\n".join(out)


def concept_map(cmid, name, title, source_cs, target_system, pairs, equivalence, description, source_vs=None, target_vs=None):
    out = [
        "Instance: " + cmid,
        "InstanceOf: ConceptMap",
        "Usage: #definition",
        f'* url = "https://ancientsky.github.io/fhir_ig_demo/ConceptMap/{cmid}"',
        f'* name = "{name}"',
        f'* title = "{fsh_str(title)}"',
        "* status = #active",
        "* experimental = false",
        f'* date = "{CODE_DATE}"',
        '* version = "0.1.0"',
        f'* description = "{fsh_str(description)}"',
    ]
    if source_vs:
        out.append(f'* sourceCanonical = "{source_vs}"')
    if target_vs:
        out.append(f'* targetCanonical = "{target_vs}"')
    groups = OrderedDict()
    for src, tgt, tdisp in pairs:
        groups.setdefault(src, []).append((tgt, tdisp))
    out += [f'* group[0].source = "{source_cs}"', f'* group[0].target = "{target_system}"']
    for i, (src, tgts) in enumerate(groups.items()):
        out.append(f"* group[0].element[{i}].code = {fsh_code(src)}")
        for j, (tgt, tdisp) in enumerate(tgts):
            out.append(f"* group[0].element[{i}].target[{j}].code = {fsh_code(tgt)}")
            if tdisp:
                out.append(f'* group[0].element[{i}].target[{j}].display = "{fsh_str(tdisp)}"')
            out.append(f"* group[0].element[{i}].target[{j}].equivalence = #{equivalence}")
    out.append("")
    return "\n".join(out)


BASE = "https://ancientsky.github.io/fhir_ig_demo"


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    rwb = openpyxl.load_workbook(REPORT_XLSX, read_only=True)
    swb = openpyxl.load_workbook(SAMPLE_XLSX, read_only=True)
    systems = []

    def simple(cid, name, title, wb, sheet, code_col=0, disp_col=1, remark_cols=(2,), skip=1, **kw):
        cs = CS(cid, name, title, sheet, **kw)
        for r in rows(wb, sheet, skip):
            if len(r) <= max(code_col, disp_col):
                continue
            remark = "；".join(r[c] for c in remark_cols if c < len(r) and r[c])
            cs.add(r[code_col], r[disp_col], remark)
        systems.append(cs)
        return cs

    # ---- 通報代碼彙整表 ----
    disease = CS(
        "cdc-disease", "CDCDisease", "CDC-NIDRS-疾病代碼", "疾病代碼",
        props=OrderedDict(
            category=("https://ancientsky.github.io/fhir_ig_demo/CodeSystem/cdc-disease#category", "string", "疾病分類(法定傳染病類別或重點監視項目)"),
            route=("https://ancientsky.github.io/fhir_ig_demo/CodeSystem/cdc-disease#route", "string", "傳染途徑"),
        ),
    )
    for r in rows(rwb, "疾病代碼"):
        if len(r) < 6:
            continue
        remark = "；".join(x for x in (r[2] if r[2] != "啟用" else "", r[3]) if x)
        disease.add(r[0], r[1], remark, category=r[4], route=r[5])
    systems.append(disease)

    determine = CS("cdc-determine-status", "CDCDetermineStatus", "CDC-NIDRS-研判結果代碼", "研判結果代碼")
    for r in rows(rwb, "研判結果代碼"):
        if r[0] == "不符合疾病分類":
            break
        if r[0] and r[1]:
            determine.add(r[0], r[1], "" if r[2] in ("保留", "新增") else r[2])
    systems.append(determine)

    determine_sub = CS(
        "cdc-determine-nonconform-reason", "CDCDetermineNonconformReason", "CDC-NIDRS-研判結果「不符合疾病分類」細項代碼", "研判結果代碼",
        description="本代碼系統為研判結果代碼 31「不符合疾病分類」之細項。",
    )
    started = False
    for r in rows(rwb, "研判結果代碼"):
        if r[0] == "不符合疾病分類":
            started = True
            continue
        if started and r[0] and r[1]:
            determine_sub.add(r[0], r[1])
    systems.append(determine_sub)

    simple("cdc-occupation", "CDCOccupation", "CDC-NIDRS-職業代碼", rwb, "職業代碼")
    simple("cdc-dense-institution", "CDCDenseInstitution", "CDC-NIDRS-人口密集機構類別代碼", rwb, "人口密集機構類別代碼")
    simple("cdc-residence-type", "CDCResidenceType", "CDC-NIDRS-非本國籍居民身分(居留)代碼", rwb, "居留代碼")
    simple("cdc-marital-status", "CDCMaritalStatus", "CDC-NIDRS-婚姻狀況代碼", rwb, "婚姻代碼")
    simple("cdc-death-reason", "CDCDeathReason", "CDC-NIDRS-個案死亡原因代碼", rwb, "死亡原因代碼")
    simple("cdc-animal", "CDCAnimal", "CDC-NIDRS-接觸動物種類代碼", rwb, "動物代碼")
    simple("cdc-symptom", "CDCSymptom", "CDC-NIDRS-症狀代碼", rwb, "症狀代碼")
    simple("cdc-china-province", "CDCChinaProvince", "CDC-NIDRS-大陸省份代碼(SARS)", rwb, "大陸省份代碼(SARS)")
    simple("cdc-hiv-risk-factor", "CDCHIVRiskFactor", "CDC-NIDRS-HIV危險因子代碼", rwb, "危險因子代碼(HIV)")
    hivdef = CS(
        "cdc-hiv-diagnosis-definition", "CDCHIVDiagnosisDefinition", "CDC-NIDRS-HIV確診定義代碼", "HIV確診定義代碼",
        props=OrderedDict(applicable=(f"{BASE}/CodeSystem/cdc-hiv-diagnosis-definition#applicable", "string", "適用對象(出生月齡)")),
        description="出生月齡定義：(通報單診斷年-個案出生年)*12 + (通報單診斷月-個案出生月) + 1。",
    )
    for r in rows(rwb, "HIV確診定義代碼"):
        hivdef.add(r[0], r[1], r[3] if len(r) > 3 else "", applicable=r[2])
    systems.append(hivdef)
    aids = CS(
        "cdc-aids-clinical-symptom", "CDCAIDSClinicalSymptom", "CDC-NIDRS-AIDS臨床症狀代碼", "AIDS臨床症狀代碼",
        props=OrderedDict(applicable=(f"{BASE}/CodeSystem/cdc-aids-clinical-symptom#applicable", "string", "適用年齡限制")),
    )
    hdr = next(rows(rwb, "AIDS臨床症狀代碼", 0))
    for r in rows(rwb, "AIDS臨床症狀代碼"):
        limits = "；".join(hdr[i] for i in (2, 3, 4) if i < len(r) and r[i])
        aids.add(r[0], r[1], applicable=limits)
    systems.append(aids)
    simple("cdc-covid-complication", "CDCCovidComplication", "CDC-NIDRS-嚴重特殊傳染性肺炎相關併發症代碼", rwb, "嚴重特殊傳染性肺炎相關併發症",
           code_col=1, disp_col=0, remark_cols=(2, 3))

    def grouped(cid, name, title, wb, sheet):
        cs = CS(cid, name, title, sheet, props=OrderedDict(item=(f"{BASE}/CodeSystem/{cid}#item", "string", "對應檢驗項目")),
                description="同一代碼可能用於多個檢驗項目，本代碼系統以「檢驗項目|代碼」作為唯一代碼，並以 item 屬性記錄對應檢驗項目。")
        for r in rows(wb, sheet):
            if r[0] and r[1] and r[2]:
                cs.add(f"{r[0]}|{r[1]}", r[2], r[3] if len(r) > 3 else "", item=r[0])
        systems.append(cs)

    grouped("cdc-tb-result", "CDCTBResult", "CDC-NIDRS-結核病檢驗結果代碼", rwb, "結核病檢驗結果代碼")
    grouped("cdc-sputum-result", "CDCSputumResult", "CDC-NIDRS-痰檢驗結果代碼(TB、MDRTB)", rwb, "痰檢驗結果代碼(TB、MDRTB)")
    simple("cdc-tb-specimen-type", "CDCTBSpecimenType", "CDC-NIDRS-結核病檢體類別代碼", rwb, "結核病檢體類別")
    simple("cdc-extrapulmonary-basis", "CDCExtrapulmonaryBasis", "CDC-NIDRS-肺外依據代碼(TB、MDRTB)", rwb, "肺外依據代碼(TB、MDRTB)")
    simple("cdc-virtual-hospital", "CDCVirtualHospital", "CDC-NIDRS-虛擬醫事機構代碼", rwb, "虛擬醫事機構代碼", remark_cols=())
    simple("cdc-quarantine-site-type", "CDCQuarantineSiteType", "CDC-NIDRS-隔離處所類型代碼", rwb, "隔離處所類型")
    qsite = CS(
        "cdc-quarantine-site", "CDCQuarantineSite", "CDC-NIDRS-隔離處所代碼", "隔離處所",
        props=OrderedDict(
            county=(f"{BASE}/CodeSystem/cdc-quarantine-site#county", "string", "處所所在縣市代碼"),
            siteType=(f"{BASE}/CodeSystem/cdc-quarantine-site#siteType", "string", "處所類型(隔離處所類型代碼)"),
        ),
    )
    for r in rows(rwb, "隔離處所"):
        qsite.add(r[0], r[1], r[4] if len(r) > 4 else "", county=r[2], siteType=r[3])
    systems.append(qsite)

    country = CS(
        "cdc-country", "CDCCountry", "CDC-NIDRS-國家代碼", "國家代碼",
        props=OrderedDict(
            alpha2=(f"{BASE}/CodeSystem/cdc-country#alpha2", "string", "ISO 3166-1 二位字母代碼"),
            alpha3=(f"{BASE}/CodeSystem/cdc-country#alpha3", "string", "ISO 3166-1 三位字母代碼"),
        ),
    )
    for r in rows(rwb, "國家代碼"):
        country.add(r[0], r[1], alpha2=r[2] if len(r) > 2 else "", alpha3=r[3] if len(r) > 3 else "")
    systems.append(country)

    county = CS("cdc-county", "CDCCounty", "CDC-NIDRS-縣市代碼", "縣市代碼(含鄉鎮市區.村里)")
    town = CS("cdc-town", "CDCTown", "CDC-NIDRS-鄉鎮市區代碼", "縣市代碼(含鄉鎮市區.村里)",
              props=OrderedDict(county=(f"{BASE}/CodeSystem/cdc-town#county", "string", "所屬縣市代碼(CDCCounty)")))
    village = CS("cdc-village", "CDCVillage", "CDC-NIDRS-村里代碼", "縣市代碼(含鄉鎮市區.村里)",
                 props=OrderedDict(town=(f"{BASE}/CodeSystem/cdc-village#town", "string", "所屬鄉鎮市區代碼(CDCTown)")))
    for r in rows(rwb, "縣市代碼(含鄉鎮市區.村里)"):
        if len(r) < 6 or not r[1]:
            continue
        remark = r[6] if len(r) > 6 else ""
        if r[1] not in county.concepts:
            county.add(r[1], r[0])
        if r[3] and r[3] not in town.concepts:
            town.add(r[3], f"{r[0]}{r[2]}", county=r[1])
        if r[5]:
            village.add(r[5], f"{r[0]}{r[2]}{r[4]}", remark, town=r[3])
    systems += [county, town, village]

    # 處理結果：只收錄可列舉之完整處理結果代碼(特殊處理代碼與各疾病通報檢核代碼)，以及項目代碼/檢核代碼之組成規則
    result = CS(
        "cdc-processing-result", "CDCProcessingResult", "CDC-NIDRS-處理結果(錯誤)代碼", "處理結果",
        description="處理結果代碼 10 碼，由「項目代碼(6碼)」加「檢核代碼(4碼)」組成，例如 IDA3020002 = 通報單-個案資料(IDA302) + 必填檢核(0002)。"
                    "本代碼系統收錄特殊處理代碼、各疾病通報檢核代碼與參考範例中之完整代碼；組成用之項目代碼與檢核代碼另見 CDCProcessingResultPart。",
    )
    result.add("SUCCESS", "成功(STATUS)")
    result.add("FAILURE", "失敗(STATUS)")
    part = CS("cdc-processing-result-part", "CDCProcessingResultPart", "CDC-NIDRS-處理結果代碼組成(項目代碼及檢核代碼)", "處理結果")
    section = None
    pending = None
    for r in rows(rwb, "處理結果", 0):
        a, b = (r + ["", ""])[:2]
        if a in ("項目代碼 (6碼)", "檢核代碼 (4碼)", "特殊處理代碼", "各疾病通報檢核代碼", "參考範例"):
            section = a
            continue
        if not a or a in ("代碼", "電子簽章與通報狀態代碼", "編碼規則"):
            continue
        if section == "項目代碼 (6碼)" and b:
            part.add(a.replace("{疾病5碼}", "{disease}"), b)
        elif section == "檢核代碼 (4碼)" and b:
            m = re.match(r"(\d{4})\s*\((.+)\)", a)
            if m:
                part.add(m.group(1), f"{m.group(2)}：{b}")
        elif section in ("特殊處理代碼", "各疾病通報檢核代碼", "參考範例"):
            if a == "說明" and pending:
                c = result.concepts[pending]
                c["comment"] = "；".join(x for x in (c["comment"], b) if x)
                continue
            if b and a not in result.concepts:
                remark = "；".join(x for x in r[2:4] if x)
                result.add(a, b, remark)
                pending = a
    systems += [result, part]

    # ---- 送驗代碼彙整表 ----
    specimen = CS("cdc-specimen-type", "CDCSpecimenType", "CDC-NIDRS-送驗檢體類別代碼", "送驗疾病與檢體類別之關聯")
    disease_specimen = []
    for r in rows(swb, "送驗疾病與檢體類別之關聯"):
        if len(r) < 4 or not r[0] or not r[2]:
            continue
        specimen.add(r[2], r[3])
        if "停用" not in (r[4] if len(r) > 4 else ""):
            disease_specimen.append((r[0], r[2], r[3]))
    systems.append(specimen)
    simple("cdc-test-disease", "CDCTestDisease", "CDC-NIDRS-疾管署之可送驗疾病代碼", swb, "疾管署之可送驗疾病代碼")
    simple("cdc-lab-unit", "CDCLabUnit", "CDC-NIDRS-指定送驗單位-疾管署單位代碼", swb, "指定送驗單位_疾管署單位代碼")
    simple("cdc-approved-lab", "CDCApprovedLab", "CDC-NIDRS-指定送驗單位-認可實驗室代碼", swb, "指定送驗單位_認可實驗室代碼", remark_cols=(), skip=2,
           description="本表異動頻繁，最新資料請至疾管署全球資訊網/應用專區/申請/傳染病認可檢驗機構(https://gov.tw/Eba)下載；此處為 v1.42 版快照。")
    simple("cdc-pathogen-method", "CDCPathogenMethod", "CDC-NIDRS-病原體檢驗方法代碼", swb, "病原體檢驗方法代碼")
    simple("cdc-tb-test-item", "CDCTBTestItem", "CDC-NIDRS-結核病送驗項目代碼", swb, "結核病送驗項目代碼")
    pathogen = CS("cdc-pathogen", "CDCPathogen", "CDC-NIDRS-病原體類別代碼(大項及細項)", "病原體類別大項、病原體類別細項及與大項之關連",
                  hierarchical=True, description="大項代碼為細項代碼之上層代碼(parent)。")
    for r in rows(swb, "病原體類別大項"):
        pathogen.add(r[0], r[1], r[2] if len(r) > 2 else "")
    for r in rows(swb, "病原體類別細項及與大項之關連"):
        if len(r) > 3 and r[2] and r[3]:
            pathogen.add(r[2], r[3], r[4] if len(r) > 4 else "", parent=r[0] if r[0] in pathogen.concepts else None)
    systems.append(pathogen)

    for cs in systems:
        (OUT / f"CodeSystem-{cs.cid}.fsh").write_text(cs.fsh(), encoding="utf-8")
        flag = f"  (略過重複代碼 {len(cs.dups)} 筆: {cs.dups[:5]})" if cs.dups else ""
        print(f"{cs.cid}: {len(cs.concepts)}{flag}")

    # ---- 概念對應 ----
    cms = []
    cms.append(concept_map(
        "cdc-country-to-iso3166", "CDCCountryToISO3166", "國家代碼對應(NIDRS國家代碼→ISO 3166-1)",
        f"{BASE}/CodeSystem/cdc-country", "urn:iso:std:iso:3166",
        [(c, v["props"]["alpha2"], "") for c, v in country.concepts.items() if v["props"].get("alpha2")],
        "equivalent", "NIDRS 國家代碼(數字碼)與 ISO 3166-1 二位字母代碼之對應。"))
    symptom_pairs = []
    for r in rows(rwb, "疾病與主要症狀關聯"):
        if len(r) > 3 and r[0] and r[2]:
            symptom_pairs.append((r[0], r[2], r[3]))
    cms.append(concept_map(
        "cdc-disease-symptom", "CDCDiseaseSymptom", "疾病與主要症狀關聯",
        f"{BASE}/CodeSystem/cdc-disease", f"{BASE}/CodeSystem/cdc-symptom", symptom_pairs, "relatedto",
        "通報疾病(DISEASE_ID)可填報之主要症狀(DISEASE_SYMPTOM)。若填寫非該疾病關聯之症狀，NIDRS 將自動改填為其他(999)，並將原症狀填入 DISEASE_SYMPTOM_OTHER。"))
    test_codes = next(cs for cs in systems if cs.cid == "cdc-test-disease").concepts
    # 關聯表中少數疾病代碼(如 010m、0703)未列於「疾管署之可送驗疾病代碼」，改以通報疾病代碼(cdc-disease)另立概念對應
    other_specimen = [x for x in disease_specimen if x[0] not in test_codes]
    disease_specimen = [x for x in disease_specimen if x[0] in test_codes]
    if other_specimen:
        cms.append(concept_map(
            "cdc-disease-specimen", "CDCDiseaseSpecimen", "送驗疾病與檢體類別之關聯(以通報疾病代碼表示者)",
            f"{BASE}/CodeSystem/cdc-disease", f"{BASE}/CodeSystem/cdc-specimen-type", other_specimen, "relatedto",
            "「送驗疾病與檢體類別之關聯」中，疾病代碼未列於「疾管署之可送驗疾病代碼」而屬通報疾病代碼者("
            + "、".join(sorted({d for d, _, _ in other_specimen})) + ")之可用檢體類別。"))
    cms.append(concept_map(
        "cdc-test-disease-specimen", "CDCTestDiseaseSpecimen", "送驗疾病與檢體類別之關聯",
        f"{BASE}/CodeSystem/cdc-test-disease", f"{BASE}/CodeSystem/cdc-specimen-type", disease_specimen, "relatedto",
        "送驗疾病(DISEASE)可使用之檢體類別(SAMPLE_TYPE)；檢體類別、屍體解剖與送驗疾病三者須符合採檢手冊之定義。"))
    (OUT / "ConceptMaps.fsh").write_text("\n".join(cms), encoding="utf-8")
    print(f"ConceptMaps: country {len(country.concepts)}, disease-symptom {len(symptom_pairs)}, disease-specimen {len(disease_specimen)}")


if __name__ == "__main__":
    main()
